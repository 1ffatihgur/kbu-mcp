"""Sayfa ve dosyalardan metin çıkarma.

Her çıkarıcı bir Extracted döner. `sections` (sayfa_no, metin) listesidir:
PDF/PPTX için gerçek sayfa/slayt numarası, diğerleri için None.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime

from selectolax.lexbor import LexborHTMLParser

from . import urls


@dataclass
class Extracted:
    text: str = ""
    sections: list[tuple[int | None, str]] = field(default_factory=list)
    title: str | None = None
    site_title: str | None = None
    links: list[tuple[str, str]] = field(default_factory=list)   # (url, link metni)
    postback_pages: list[tuple[str, str]] = field(default_factory=list)  # (__EVENTTARGET, __EVENTARGUMENT)
    content_date: str | None = None
    pages: int | None = None
    needs_ocr: bool = False


# ---------------------------------------------------------------- tarihler

_MONTHS = {
    "ocak": 1, "şubat": 2, "subat": 2, "mart": 3, "nisan": 4, "mayıs": 5, "mayis": 5, "haziran": 6,
    "temmuz": 7, "ağustos": 8, "agustos": 8, "eylül": 9, "eylul": 9, "ekim": 10, "kasım": 11,
    "kasim": 11, "aralık": 12, "aralik": 12,
}
_DATE_NUM = re.compile(r"\b(\d{1,2})[./](\d{1,2})[./](20\d{2})\b")
_DATE_TXT = re.compile(r"\b(\d{1,2})\s+(" + "|".join(_MONTHS) + r")\s+(20\d{2})\b", re.IGNORECASE)


def find_date(text: str) -> str | None:
    """Metindeki ilk makul tarihi ISO biçiminde döner."""
    head = text[:3000]
    cands = []
    for m in _DATE_NUM.finditer(head):
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        cands.append((m.start(), d, mo, y))
    for m in _DATE_TXT.finditer(head):
        mo = _MONTHS.get(m.group(2).lower().replace("İ", "i"))
        if mo:
            cands.append((m.start(), int(m.group(1)), mo, int(m.group(3))))
    for _, d, mo, y in sorted(cands):
        try:
            return datetime(y, mo, d).date().isoformat()
        except ValueError:
            continue
    return None


def clean_text(s: str) -> str:
    s = s.replace("\xa0", " ").replace("\u200b", "")
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    s = re.sub(r" *\n *", "\n", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


# ---------------------------------------------------------------- HTML

# Site şablonlarında asıl içeriğin bulunduğu alanlar (öncelik sırasıyla).
CONTENT_SELECTORS = [
    "#txtIcerik", "#lblPostIcerik_Sayfa", "[id$=lblIcerik]", "[id$=txtIcerik]", "[id*=Icerik]",
    ".entry-content", "article", "main", "#MainBlock", "#content", ".content",
]
# Şablonun her sayfada tekrar eden kısımları.
BOILERPLATE_SELECTORS = [
    "script", "style", "noscript", "iframe", "svg", "form input", "nav", "header .menu",
    "#main-nav", "#main_nav", ".main-navigation", "#site-navigation", "#tophead", "#news-ticker",
    "#duyuruPage", "#duyuruModal", "#myModal", "#popupClass", ".modal", "#footerGrup",
    "#footer-widgets", ".footer-widgets", ".skip-link", ".screen-reader-text", ".breadcrumb",
]
FOOTER_CONTACT_SELECTORS = ["#Master_Label_Footer_Iletisim", "footer .iletisim", "#iletisim", ".site-info", "footer"]


def _node_text(node) -> str:
    """Blok elemanlarında satır sonu koruyarak metin; tablolar satır satır."""
    for br in node.css("br"):
        br.replace_with("\n")
    for tr in node.css("tr"):
        cells = [c.text(separator=" ", strip=True) for c in tr.css("td, th")]
        tr.replace_with("\n" + " | ".join(c for c in cells if c) + "\n")
    for blk in node.css("p, div, li, h1, h2, h3, h4, h5, h6, section, article, ul, ol, table, dd, dt"):
        blk.insert_before("\n")
        blk.insert_after("\n")
    for li in node.css("li"):
        li.insert_before("• ")
    return clean_text(node.text(separator="", strip=False))


def _postback_pages(html: str) -> list[tuple[str, str]]:
    out = []
    for m in re.finditer(r"__doPostBack\(&#39;([^&]+?)&#39;,&#39;(Page\$\d+|Page\$Next|Page\$Last)&#39;\)|"
                         r"__doPostBack\('([^']+?)','(Page\$\d+|Page\$Next|Page\$Last)'\)", html):
        target = m.group(1) or m.group(3)
        arg = m.group(2) or m.group(4)
        if (target, arg) not in out:
            out.append((target, arg))
    return out


def extract_html(html: str, url: str) -> Extracted:
    tree = LexborHTMLParser(html)
    ex = Extracted()
    t = tree.css_first("title")
    ex.site_title = clean_text(t.text()) if t else None

    for a in tree.css("a[href], area[href], iframe[src], frame[src]"):
        href = a.attributes.get("href") or a.attributes.get("src") or ""
        u = urls.normalize(href, url)
        if u:
            ex.links.append((u, clean_text(a.text(separator=" "))[:200]))
    # window.open('...') / location.href='...' içindeki linkler
    # (Bologna menüsü: menu_close(this,'progCourses.aspx?...'))
    for m in re.finditer(r"(?:window\.open|location\.href\s*=|menu_close\(this,)\s*\(?\s*['\"]([^'\"]+)['\"]", html):
        u = urls.normalize(m.group(1), url)
        if u:
            ex.links.append((u, ""))
    ex.postback_pages = _postback_pages(html)

    footer_contact = ""
    for sel in FOOTER_CONTACT_SELECTORS:
        n = tree.css_first(sel)
        if n:
            txt = clean_text(n.text(separator="\n"))
            if re.search(r"tel|telefon|e-?posta|adres", txt, re.I) and len(txt) < 1500:
                footer_contact = txt
                break

    body = tree.body
    if body is None:
        return ex
    for sel in BOILERPLATE_SELECTORS:
        for n in body.css(sel):
            n.decompose()

    content = None
    for sel in CONTENT_SELECTORS:
        cands = [n for n in body.css(sel) if len(n.text(strip=True)) > 40]
        if cands:
            content = max(cands, key=lambda n: len(n.text(strip=True)))
            break

    path = urls.urlsplit(url).path.lower().rstrip("/")
    # Site / bölüm ana sayfaları (ör. /bilgisayar): duyuru blokları ve iletişim bilgisiyle birlikte tamamı alınır.
    is_home = path in ("", "/index.aspx", "/default.aspx", "/index.html", "/anasayfa", "/en") or bool(
        re.fullmatch(r"/[a-z]+(/index\.aspx)?", path)
    )

    if content is not None and not is_home:
        heading = None
        for h in body.css("h1, h2, h3, .entry-title, .page-title, [id$=lblBaslik], [id*=Baslik]"):
            ht = clean_text(h.text(separator=" "))
            if 3 < len(ht) < 200 and ht != ex.site_title:
                heading = ht
                break
        ex.title = heading
        text = _node_text(content)
        if heading and heading not in text[:300]:
            text = heading + "\n\n" + text
    else:
        text = _node_text(body)

    if footer_contact and (is_home or re.search(r"iletisim|iletişim|contact", url + (ex.title or ""), re.I)):
        text += "\n\nİletişim bilgileri:\n" + footer_contact

    ex.text = clean_text(text)
    if not ex.title:
        first = ex.text.split("\n", 1)[0].strip()
        ex.title = first[:150] if first else ex.site_title
    ex.sections = [(None, ex.text)]
    ex.content_date = find_date(ex.text)
    return ex


# ---------------------------------------------------------------- PDF

def extract_pdf(data: bytes) -> Extracted:
    import pymupdf

    ex = Extracted()
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        ex.pages = doc.page_count
        meta_title = (doc.metadata or {}).get("title") or ""
        total = 0
        image_pages = 0
        for i, page in enumerate(doc, start=1):
            txt = clean_text(page.get_text("text", sort=True))
            total += len(txt)
            if len(txt) < 30 and page.get_images():
                image_pages += 1
            if txt:
                ex.sections.append((i, txt))
            for link in page.get_links():
                uri = link.get("uri")
                u = urls.normalize(uri) if uri else None
                if u:
                    ex.links.append((u, ""))
        ex.needs_ocr = ex.pages > 0 and (image_pages / ex.pages > 0.5 or total < 40 * ex.pages)
    ex.text = "\n\n".join(t for _, t in ex.sections)
    if meta_title and len(meta_title) > 4 and not re.match(r"^(microsoft|untitled|başlıksız|document|belge)", meta_title, re.I):
        ex.title = meta_title.strip()
    elif ex.text:
        ex.title = ex.text.strip().split("\n", 1)[0][:150]
    ex.content_date = find_date(ex.text)
    return ex


_MONTH_WORD = re.compile(r"^(OCAK|ŞUBAT|MART|NİSAN|MAYIS|HAZİRAN|TEMMUZ|AĞUSTOS|EYLÜL|EKİM|KASIM|ARALIK)$", re.I)


def pdf_date_grid(data: bytes) -> list[tuple[str, list[str]]]:
    """Takvim düzenli PDF'lerde (ör. aylık yemek listesi) her tarih başlığının altındaki metni o güne atar.

    Düz metin çıkarımı sütunları karıştırır, tablo algılama da birleşik hücrelerde yanılır.
    Burada kelime koordinatları kullanılır: bir kelime, aynı satır bandındaki x-merkezi en yakın
    tarih başlığına aittir.
    """
    import pymupdf

    out: list[tuple[str, list[str]]] = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            words = page.get_text("words")
            heads = []  # (x_merkez, y0, y1, etiket, kelime_indeksleri)
            for i in range(len(words) - 2):
                w = words[i]
                if re.fullmatch(r"\d{1,2}", w[4]) and _MONTH_WORD.match(words[i + 1][4]) and re.fullmatch(r"20\d{2}", words[i + 2][4]):
                    idx = [i, i + 1, i + 2]
                    label = f"{w[4]} {words[i + 1][4]} {words[i + 2][4]}"
                    if i + 3 < len(words) and re.fullmatch(r"[A-ZÇĞİÖŞÜ]{4,}", words[i + 3][4]) and abs(words[i + 3][1] - w[1]) < 3:
                        label += " " + words[i + 3][4]
                        idx.append(i + 3)
                    x1 = words[idx[-1]][2]
                    heads.append(((w[0] + x1) / 2, w[1], w[3], label, set(idx)))
            if not heads:
                continue
            header_idx = set().union(*(h[4] for h in heads))
            row_ys = sorted({round(h[1]) for h in heads})
            for h in heads:
                row = round(h[1])
                nxt = [y for y in row_ys if y > row + 2]
                band_end = nxt[0] if nxt else page.rect.height
                same_row = [x for x in heads if abs(round(x[1]) - row) <= 2]
                cell = []
                for j, w in enumerate(words):
                    if j in header_idx or not (h[2] - 1 <= w[1] < band_end - 1):
                        continue
                    xc = (w[0] + w[2]) / 2
                    if min(same_row, key=lambda x: abs(x[0] - xc)) is h:
                        cell.append(w)
                lines: dict[int, list] = {}
                for w in sorted(cell, key=lambda w: (w[1], w[0])):
                    key = next((k for k in lines if abs(k - w[1]) < 3), round(w[1]))
                    lines.setdefault(key, []).append(w[4])
                items = []
                for _, v in sorted(lines.items()):
                    line = " ".join(v)
                    if re.search(r"[a-zçğıöşü]", line):  # menü büyük harf; altındaki imza bloğu değil
                        break
                    if len(line) > 2:
                        items.append(line)
                out.append((h[3], items))
    return out


# ---------------------------------------------------------------- Office

def extract_docx(data: bytes) -> Extracted:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = []
    # Paragraf ve tabloları belgede geçtikleri sırayla oku.
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in block.iter() if n.tag.endswith("}t"))
            if t.strip():
                parts.append(t)
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iterchildren():
                        if cell.tag.endswith("}tc"):
                            ct = " ".join((n.text or "") for n in cell.iter() if n.tag.endswith("}t")).strip()
                            if ct and (not cells or cells[-1] != ct):
                                cells.append(ct)
                    if cells:
                        parts.append(" | ".join(cells))
    ex = Extracted(text=clean_text("\n".join(parts)))
    ex.sections = [(None, ex.text)]
    ex.title = (d.core_properties.title or "").strip() or (ex.text.split("\n", 1)[0][:150] if ex.text else None)
    ex.content_date = find_date(ex.text)
    return ex


def extract_xlsx(data: bytes) -> Extracted:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ex = Extracted()
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
            if len(rows) > 20000:
                break
        if rows:
            ex.sections.append((None, f"[Sayfa: {ws.title}]\n" + "\n".join(rows)))
    wb.close()
    ex.text = clean_text("\n\n".join(t for _, t in ex.sections))
    ex.title = ex.text.split("\n", 2)[1][:150] if ex.text.count("\n") >= 1 else None
    ex.content_date = find_date(ex.text)
    return ex


def extract_pptx(data: bytes) -> Extracted:
    import pptx

    prs = pptx.Presentation(io.BytesIO(data))
    ex = Extracted(pages=len(prs.slides))
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells if c.text))
        t = clean_text("\n".join(parts))
        if t:
            ex.sections.append((i, t))
    ex.text = "\n\n".join(t for _, t in ex.sections)
    ex.title = ex.text.split("\n", 1)[0][:150] if ex.text else None
    ex.content_date = find_date(ex.text)
    return ex


def extract_plain(data: bytes) -> Extracted:
    for enc in ("utf-8", "cp1254", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    ex = Extracted(text=clean_text(text))
    ex.sections = [(None, ex.text)]
    ex.title = ex.text.split("\n", 1)[0][:150] if ex.text else None
    return ex


def extract_odf(data: bytes) -> Extracted:
    """ODT/ODS: content.xml içindeki metin."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml").decode("utf-8", "ignore")
    xml = re.sub(r"</text:p>|</text:h>|</table:table-row>", "\n", xml)
    xml = re.sub(r"</table:table-cell>", " | ", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    return extract_plain(text.encode("utf-8"))


def extract_rtf(data: bytes) -> Extracted:
    s = data.decode("latin-1", "ignore")
    s = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1254", "ignore"), s)
    s = re.sub(r"\\par[d]?", "\n", s)
    s = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", s)
    return extract_plain(s.encode("utf-8"))


def extract_zip(data: bytes) -> Extracted:
    """ZIP içindeki belgelerin metinlerini tek belge altında birleştirir."""
    ex = Extracted()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist()[:200]:
            if info.is_dir() or info.file_size > 30 * 1024 * 1024:
                continue
            name = info.filename
            try:
                name = name.encode("cp437").decode("cp857")  # Türkçe Windows ZIP adları
            except (UnicodeEncodeError, UnicodeDecodeError):
                pass
            kind = urls.kind_from_url("x://h/" + name)
            if not kind or kind in ("zip",):
                continue
            try:
                sub = extract_bytes(z.read(info), kind)
            except Exception:
                continue
            if sub and sub.text:
                ex.sections.append((None, f"[Dosya: {name}]\n{sub.text}"))
    ex.text = "\n\n".join(t for _, t in ex.sections)
    ex.title = None
    return ex


def extract_bytes(data: bytes, kind: str) -> Extracted | None:
    if kind == "pdf":
        return extract_pdf(data)
    if kind == "docx":
        return extract_docx(data)
    if kind == "xlsx":
        return extract_xlsx(data)
    if kind == "pptx":
        return extract_pptx(data)
    if kind in ("txt", "csv"):
        return extract_plain(data)
    if kind in ("odt", "ods"):
        return extract_odf(data)
    if kind == "rtf":
        return extract_rtf(data)
    if kind == "zip":
        return extract_zip(data)
    if kind in ("doc", "xls", "ppt"):
        return extract_legacy_office(data, kind)
    return None


def extract_legacy_office(data: bytes, kind: str) -> Extracted | None:
    """Eski .doc/.xls/.ppt: bazıları aslında yeni format (yanlış uzantı); değilse kaba metin çıkarımı."""
    if data[:2] == b"PK":
        return extract_bytes(data, {"doc": "docx", "xls": "xlsx", "ppt": "pptx"}[kind])
    if data[:5] == b"%PDF-":
        return extract_pdf(data)
    if data[:5] in (b"<html", b"<HTML", b"<!DOC") or data[:20].lstrip().lower().startswith(b"<"):
        return extract_html(data.decode("utf-8", "ignore"), "https://www.karabuk.edu.tr/")
    # OLE ikili biçim: UTF-16LE metin parçalarını topla (Word metni çoğunlukla böyle saklar).
    runs = re.findall(rb"(?:[\x20-\x7e\xc0-\xff\x00-\x01][\x00\x01]){6,}", data)
    text = "\n".join(r.decode("utf-16le", "ignore") for r in runs)
    text = re.sub(r"[^\w\s.,;:!?()/%&@+\-–'\"|•]", " ", text)
    ex = extract_plain(text.encode("utf-8"))
    ex.needs_ocr = len(ex.text) < 200
    return ex


def sniff_kind(data: bytes, content_type: str, url_kind: str | None) -> str | None:
    ct = (content_type or "").lower()
    if data[:5] == b"%PDF-":
        return "pdf"
    if "html" in ct:
        return "html"
    if url_kind:
        return url_kind
    if "pdf" in ct:
        return "pdf"
    if "wordprocessingml" in ct:
        return "docx"
    if "spreadsheetml" in ct:
        return "xlsx"
    if "presentationml" in ct:
        return "pptx"
    if "msword" in ct:
        return "doc"
    if "ms-excel" in ct:
        return "xls"
    if "ms-powerpoint" in ct:
        return "ppt"
    if ct.startswith("text/plain"):
        return "txt"
    return None
